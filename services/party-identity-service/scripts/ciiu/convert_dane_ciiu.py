"""Convierte la estructura detallada oficial de la CIIU del DANE al CSV del `seed` (R-10, FR-056).

Fuera del paquete del servicio: solo se ejecuta al actualizar el snapshot, nunca en la carga.

Entrada: la hoja de cálculo oficial guardada sin modificar en
`src/party_identity_service/seed/data/ciiu_rev5_ac/source/` (hoja "CIIU 5 A.C.", columnas
`Categoría`, código incluida la sección, código y título).

Salida: `ciiu_rev5_ac.csv` (UTF-8, saltos de línea LF) con columnas `level`, `code`, `name`,
`parent_code`, en el orden del archivo oficial. El padre se deriva del propio código:

- sección (letra) → sin padre;
- división (2 dígitos) → la sección del código con sección (`A01` → `A`);
- grupo (3 dígitos) → división de sus 2 primeros dígitos;
- clase (4 dígitos) → grupo de sus 3 primeros dígitos.

Imprime los SHA-256 del original y del CSV y los conteos por nivel para copiarlos en `SOURCE.md`.

Uso (desde `services/party-identity-service/`, con el extra `tools` instalado:
`pip install -e ".[tools]"`):

    python scripts/ciiu/convert_dane_ciiu.py

El resultado es reproducible: el mismo archivo oficial produce siempre los mismos bytes.
"""

from __future__ import annotations

import csv
import hashlib
import sys
from collections import Counter
from collections.abc import Iterable, Sequence
from pathlib import Path

SERVICE_DIR = Path(__file__).resolve().parents[2]
SNAPSHOT_DIR = SERVICE_DIR / "src" / "party_identity_service" / "seed" / "data" / "ciiu_rev5_ac"
SOURCE_FILE = SNAPSHOT_DIR / "source" / "CIIU_5_AC_Estructura_detallada.xlsx"
CSV_FILE = SNAPSHOT_DIR / "ciiu_rev5_ac.csv"
SHEET = "CIIU 5 A.C."

COLUMNS = ("level", "code", "name", "parent_code")
LEVELS = {"Sección": ("SECTION", 1), "División": ("DIVISION", 2), "Grupo": ("GROUP", 3), "Clase": ("CLASS", 4)}


def derive_rows(rows: Iterable[Sequence[object]]) -> list[dict[str, str]]:
    """Filas del archivo oficial → filas del CSV con nivel y código padre.

    Ignora todo lo anterior al encabezado (`Categoría`) y las filas vacías.

    Raises:
        ValueError: Encabezado ausente, categoría desconocida, código con longitud o sección
            incoherentes, título vacío, código duplicado o padre inexistente.
    """
    result: list[dict[str, str]] = []
    seen: set[str] = set()
    header_found = False

    for number, raw in enumerate(rows, start=1):
        cells = [str(cell).strip() if cell is not None else "" for cell in raw[:4]]
        cells += [""] * (4 - len(cells))
        category, full_code, code, name = cells

        if not header_found:
            header_found = category == "Categoría"
            continue
        if not any(cells):
            continue

        if category not in LEVELS:
            raise ValueError(f"Fila {number}: Categoría desconocida {category!r}.")
        level, length = LEVELS[category]

        if len(code) != length:
            raise ValueError(f"Fila {number}: el código {code!r} no tiene la longitud de {category} ({length}).")
        section = full_code[:-length] if level != "SECTION" else ""
        if full_code != f"{section}{code}" or (level != "SECTION" and len(section) != 1):
            raise ValueError(f"Fila {number}: el código con sección {full_code!r} no coincide con {code!r}.")
        if not name:
            raise ValueError(f"Fila {number}: título vacío para {code!r}.")
        if code in seen:
            raise ValueError(f"Fila {number}: código duplicado {code!r}.")

        parent = {"SECTION": "", "DIVISION": section, "GROUP": code[:2], "CLASS": code[:3]}[level]
        if parent and parent not in seen:
            raise ValueError(f"Fila {number}: el padre {parent!r} de {code!r} no aparece antes en el archivo.")

        seen.add(code)
        result.append({"level": level, "code": code, "name": name, "parent_code": parent})

    if not header_found:
        raise ValueError("No se encontró el encabezado 'Categoría' del archivo oficial.")

    return result


def read_workbook(path: Path) -> list[tuple[object, ...]]:
    """Lee la hoja oficial (requiere `openpyxl`, extra `tools`)."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook[SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()


def write_csv(rows: list[dict[str, str]], path: Path) -> None:
    """Escribe el CSV en UTF-8 con LF para que su SHA-256 sea estable."""
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    rows = derive_rows(read_workbook(SOURCE_FILE))
    write_csv(rows, CSV_FILE)

    counts = Counter(row["level"] for row in rows)
    print(f"SHA-256 original ({SOURCE_FILE.name}): {sha256(SOURCE_FILE)}")
    print(f"SHA-256 CSV ({CSV_FILE.name}): {sha256(CSV_FILE)}")
    for level, _ in LEVELS.values():
        print(f"{level}: {counts[level]}")
    print(f"Total: {len(rows)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
